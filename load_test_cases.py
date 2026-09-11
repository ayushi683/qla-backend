import os
from datetime import datetime, timezone
from decimal import Decimal

from app.database import SessionLocal, init_db
from app.models.party import Party
from app.models.inquiry_case import InquiryCase
from app.models.extract import ExtractedLineItem
from app.models.email import EmailThread, EmailMessage
from app.models.document import InquiryDocument

TEST_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_enquiries")
ENQUIRY_DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "instance", "enquiry_docs")

REF_PREFIX = "TEST-"


def extract_text_snippet(filepath, max_chars=800):
    ext = filepath.lower().rsplit(".", 1)[-1] if "." in filepath else ""
    text = ""
    try:
        if ext == "pdf":
            import pypdf
            reader = pypdf.PdfReader(filepath)
            for page in reader.pages[:2]:
                text += page.extract_text() or ""
        elif ext in ("xlsx", "xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            rows_text = []
            for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
                rows_text.append(" ".join(str(c) for c in row if c is not None))
            text = "\n".join(rows_text)
        elif ext == "docx":
            import docx
            d = docx.Document(filepath)
            text = "\n".join(p.text for p in d.paragraphs[:20])
    except Exception as e:
        print("  (could not extract text from " + os.path.basename(filepath) + ": " + str(e) + ")")

    text = text.strip()
    if not text:
        return None
    return text[:max_chars]


def content_type_for(filepath):
    ext = filepath.lower().rsplit(".", 1)[-1] if "." in filepath else ""
    return {
        "pdf": "application/pdf",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xls": "application/vnd.ms-excel",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "doc": "application/msword",
    }.get(ext, "application/octet-stream")


def pick_best_file_for_text(files_in_folder):
    """Prefer the covering enquiry email PDF (usually named *-ENQ.pdf or
    similar) for the description text — fall back to the first PDF,
    then the first file of any kind."""
    pdfs = [f for f in files_in_folder if f.lower().endswith(".pdf")]
    enq_pdfs = [f for f in pdfs if "enq" in f.lower()]
    if enq_pdfs:
        return enq_pdfs[0]
    if pdfs:
        return pdfs[0]
    return files_in_folder[0] if files_in_folder else None


def main():
    if not os.path.isdir(TEST_DIR):
        print("Folder not found: " + TEST_DIR)
        print("Create it and put per-enquiry subfolders inside, then run this again.")
        return

    subfolders = sorted([
        d for d in os.listdir(TEST_DIR)
        if os.path.isdir(os.path.join(TEST_DIR, d))
    ])
    if not subfolders:
        print("No subfolders found in " + TEST_DIR)
        print("Expected structure: test_enquiries/<enquiry-name>/file1.pdf, file2.xlsx, ...")
        return

    init_db()
    db = SessionLocal()

    # Wipe out cases created by a previous run of this script.
    old_cases = db.query(InquiryCase).filter(InquiryCase.internal_ref.like(REF_PREFIX + "%")).all()
    for oc in old_cases:
        db.query(ExtractedLineItem).filter_by(case_id=oc.case_id).delete()
        db.query(InquiryDocument).filter_by(case_id=oc.case_id).delete()
        threads = db.query(EmailThread).filter_by(case_id=oc.case_id).all()
        for t in threads:
            db.query(EmailMessage).filter_by(thread_id=t.thread_id).delete()
            db.delete(t)
        db.delete(oc)
    db.commit()
    if old_cases:
        print("Cleared " + str(len(old_cases)) + " cases from a previous run.")

    for folder_name in subfolders:
        folder_path = os.path.join(TEST_DIR, folder_name)
        files_in_folder = sorted([
            f for f in os.listdir(folder_path)
            if os.path.isfile(os.path.join(folder_path, f))
        ])
        if not files_in_folder:
            print("Skipping empty folder: " + folder_name)
            continue

        internal_ref = REF_PREFIX + folder_name

        best_file = pick_best_file_for_text(files_in_folder)
        snippet = extract_text_snippet(os.path.join(folder_path, best_file)) if best_file else None

        subject = "Enquiry — " + folder_name
        body_text = snippet or ("(Could not auto-extract text — see attached files.)")
        customer_name = "Test Customer — " + folder_name

        customer = Party(party_type="CUSTOMER", display_name=customer_name, email=None)
        db.add(customer)
        db.flush()

        case = InquiryCase(
            internal_ref=internal_ref, status="RECEIVED", source_type="NORMAL",
            project_name=folder_name,
            enq_received_at=datetime.now(timezone.utc),
            customer_party_id=customer.party_id,
        )
        db.add(case)
        db.flush()

        thread = EmailThread(
            case_id=case.case_id, mailbox="enquiries@punetechtrol.com",
            subject_normalized=subject, first_message_at=datetime.now(timezone.utc),
            last_message_at=datetime.now(timezone.utc),
        )
        db.add(thread)
        db.flush()

        message = EmailMessage(
            thread_id=thread.thread_id, case_id=case.case_id, direction="INBOUND",
            sender_email="test@example.com", subject=subject, body_text=body_text,
            received_at=datetime.now(timezone.utc), has_attachments=True,
        )
        db.add(message)
        db.flush()

        case_dir = os.path.join(ENQUIRY_DOCS_DIR, str(case.case_id))
        os.makedirs(case_dir, exist_ok=True)

        for filename in files_in_folder:
            src_path = os.path.join(folder_path, filename)
            dest_path = os.path.join(case_dir, filename)
            with open(src_path, "rb") as src_f:
                content = src_f.read()
            with open(dest_path, "wb") as dst_f:
                dst_f.write(content)

            db.add(InquiryDocument(
                case_id=case.case_id, message_id=message.message_id, file_name=filename,
                relative_path=os.path.join("enquiry_docs", str(case.case_id), filename),
                blob_uri=os.path.join("enquiry_docs", str(case.case_id), filename),
                doc_role="ENQUIRY", content_type=content_type_for(filename), size_bytes=len(content),
            ))

        line = ExtractedLineItem(
            case_id=case.case_id, line_no=1,
            description=(snippet[:200] if snippet else folder_name),
            qty=Decimal("1"), uom="NOS",
        )
        db.add(line)

        db.commit()
        print("Created " + internal_ref + " with " + str(len(files_in_folder)) + " file(s): " + ", ".join(files_in_folder))

    db.close()
    print("Done.")


if __name__ == "__main__":
    main()